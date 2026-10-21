import copy
import json
import random

import pytest
from openpyxl import load_workbook

from kkai.code_ai import parse_response_body, request_body, coder_materials, run_code
from kkai.codebook import check_flags, flatten, load_schema, random_output, validate_output
from kkai.generate import run_generate
from kkai.ingest import run_ingest
from kkai.plan import write_plan
from kkai.prepare import read_key, run_prepare
from kkai.sheets import run_sheet
from kkai.util import KkaiError, read_csv


@pytest.fixture
def schema(cfg):
    return load_schema(str(cfg.path(cfg.raw["coder"]["schema"])))


def test_random_output_is_valid(schema):
    rng = random.Random(1)
    for _ in range(200):
        out = random_output(schema, rng)
        assert validate_output(schema, out) == []
        assert check_flags(out)["flags_missing"] == ""


def test_evidence_must_cite_existing_observation(schema):
    out = random_output(schema, random.Random(2))
    out["medium"]["evidence"] = ["o99"]
    assert any("o99" in e for e in validate_output(schema, out))


def test_schema_violation_detected(schema):
    out = random_output(schema, random.Random(3))
    out["temporal_frame"]["value"] = "medieval"
    assert validate_output(schema, out)


def test_program_recomputes_flags(schema):
    out = random_output(schema, random.Random(4))
    out["religion"]["form"] = "building"
    out["review_reasons"] = []
    out["needs_human_review"] = False
    f = check_flags(out)
    assert "religion_present" in f["flags_missing"]
    assert f["needs_review_mismatch"] and f["needs_review_program"]


def test_identity_terms_flagged(schema):
    out = random_output(schema, random.Random(5))
    out["observations"][0]["text"] = "a Hasidic man walking"
    assert check_flags(out)["identity_terms"] == "hasidic"


def test_temporal_rule(schema):
    out = random_output(schema, random.Random(6))
    out["temporal_frame"].update(value="historical", contemporary_cue_types=["mobile_device", "modern_tram"],
                                 historical_cue_types=[])
    assert check_flags(out)["temporal_rule_mismatch"]


def test_flatten_joins_lists(schema):
    out = random_output(schema, random.Random(7))
    out["figure_roles"] = [{"role": "sightseeing", "evidence": ["o1"]}, {"role": "everyday_activity", "evidence": ["o1"]},
                           {"role": "sightseeing", "evidence": ["o2"]}]
    assert flatten(out)["roles"] == "everyday_activity;sightseeing"


def test_request_body_is_blind(cfg):
    mats = coder_materials(cfg)
    body = request_body(cfg, mats, b"\x89PNGfake")
    text = json.dumps(body)
    for leak in ("W-Z", "multicultural", "Kazimierz in", "gpt-image", "nano-banana", "img_"):
        assert leak not in text.replace(mats["instruction"], "")
    assert body["text"]["format"]["strict"] is True
    user = body["input"][1]["content"]
    assert user[0]["text"] == "Code this image according to the instruction."
    assert user[1]["image_url"].startswith("data:image/png;base64,")


def test_parse_response_body():
    ok = {"status": "completed", "model": "m",
          "output": [{"type": "reasoning"}, {"type": "message", "content": [{"type": "output_text", "text": "{\"a\": 1}"}]}]}
    assert parse_response_body(ok) == ({"a": 1}, None)
    ref = {"status": "completed", "output": [{"type": "message", "content": [{"type": "refusal", "refusal": "no"}]}]}
    assert parse_response_body(ref)[1].startswith("refusal")
    inc = {"status": "incomplete", "incomplete_details": {"reason": "max_output_tokens"}}
    assert "incomplete" in parse_response_body(inc)[1]


def _pilot(cfg):
    write_plan(cfg, "pilot")
    run_generate(cfg, "pilot", assume_yes=True)
    run_prepare(cfg, "pilot")


def test_dry_run_coding_and_resume(cfg):
    _pilot(cfg)
    res = run_code(cfg, "pilot", assume_yes=True)
    n = len(read_key(cfg))
    assert res["status"] == "done" and res["rows"] == n and res["invalid"] == 0
    again = run_code(cfg, "pilot", assume_yes=True)
    assert again["status"] == "done"
    assert len(read_csv(cfg.data_dir / "ai" / "ai_codes.csv")) == n
    new = run_code(cfg, "pilot", assume_yes=True, new_run=True)
    assert new["run_id"] == "pilot-r2" and new["rows"] == 2 * n


class FakeBatchClient:
    """Minimal stand-in for the OpenAI files/batches API."""

    def __init__(self, schema):
        self.schema = schema
        self.uploaded = {}
        self.status = "in_progress"
        outer = self

        class Files:
            def create(self, file, purpose):
                fid = f"file_{len(outer.uploaded)}"
                outer.uploaded[fid] = file.read().decode()
                return type("F", (), {"id": fid})

            def content(self, fid):
                return type("C", (), {"text": outer.outputs[fid]})

        class Batches:
            def create(self, input_file_id, endpoint, completion_window, metadata):
                assert endpoint == "/v1/responses" and completion_window == "24h"
                outer.batch_input = input_file_id
                return type("B", (), {"id": "batch_1"})

            def retrieve(self, bid):
                outer.outputs = {}
                lines = []
                rng = random.Random(0)
                for line in outer.uploaded[outer.batch_input].splitlines():
                    req = json.loads(line)
                    out = random_output(outer.schema, rng)
                    body = {"status": "completed", "model": "gpt-test", "output": [
                        {"type": "message", "content": [{"type": "output_text", "text": json.dumps(out)}]}]}
                    lines.append(json.dumps({"id": "x", "custom_id": req["custom_id"],
                                             "response": {"status_code": 200, "request_id": "r", "body": body},
                                             "error": None}))
                outer.outputs["out_1"] = "\n".join(lines)
                return type("S", (), {"status": outer.status, "output_file_id": "out_1" if outer.status == "completed" else None,
                                      "error_file_id": None, "request_counts": {}})

        self.files, self.batches = Files(), Batches()


def test_batch_flow(cfg, schema):
    _pilot(cfg)
    cfg.dry_run_backup = cfg.dry_run
    cfg.raw["coder"]["mode"] = "batch"
    client = FakeBatchClient(schema)
    # Force batch mode while keeping data in the dry-run folder.
    cfg.raw["data_dir"] = str(cfg.data_dir)
    cfg.dry_run = False
    first = run_code(cfg, "pilot", assume_yes=True, client=client)
    assert first["status"] == "submitted"
    assert run_code(cfg, "pilot", assume_yes=True, client=client)["status"] == "pending"
    client.status = "completed"
    done = run_code(cfg, "pilot", assume_yes=True, client=client)
    assert done["status"] == "done" and done["invalid"] == 0
    assert done["rows"] == len(read_key(cfg))


def test_sheet_and_ingest(cfg):
    _pilot(cfg)
    xlsx, page = run_sheet(cfg, "pilot")
    with pytest.raises(KkaiError):
        run_sheet(cfg, "pilot")
    html_text = page.read_text(encoding="utf-8")
    key = read_key(cfg)
    assert all(k["condition_code"] not in html_text for k in key)
    wb = load_workbook(xlsx)
    ws = wb["codes"]
    header = [c.value for c in ws[1]]
    assert header[0] == "image_id" and "temporal_frame" in header
    assert not any(k["condition_code"] in str(c.value) for row in ws.iter_rows() for c in row for k in key)
    col = {h: i + 1 for i, h in enumerate(header)}
    for r in range(2, ws.max_row + 1):
        for name, val in [("medium", "photo_colour"), ("temporal_frame", "contemporary"), ("mode_dominant", "heritage"),
                          ("religion_form", "none"), ("place", "landmark"), ("roles", "sightseeing; everyday_activity")]:
            ws.cell(r, col[name], val)
    ws.cell(2, col["landmarks"], "wawel;atlantis")
    filled = xlsx.with_name("pilot_filled.xlsx")
    wb.save(filled)
    with pytest.raises(KkaiError):
        run_ingest(cfg, filled, "pilot", "PB")
    ws.cell(2, col["landmarks"], "wawel")
    wb.save(filled)
    assert run_ingest(cfg, filled, "pilot", "PB") == len(key)
    rows = read_csv(cfg.data_dir / "human" / "human_codes.csv")
    assert rows[0]["roles"] in ("everyday_activity;sightseeing",)
    with pytest.raises(KkaiError):
        run_ingest(cfg, filled, "pilot", "PB")
    assert run_ingest(cfg, filled, "pilot", "PB", round_no=2) == len(key)
