"""Review of the remaining main-stage images after gate 3, and the final codes.

decisions.csv (data/human/) holds your gate-3 decision per variable:
review_flagged, review_all or manual_all. The review sheet keeps the model
codes in their own columns; your corrections go into separate columns.
"""

from __future__ import annotations

import csv
import html
import os
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill

from .codebook import SHEET_COLUMNS, allowed_values, load_schema
from .config import Config
from .prepare import read_key
from .util import KkaiError, read_csv, sha256_file, utc_now, write_csv_once, write_new_text

DECISIONS = ("review_flagged", "review_all", "manual_all")
# Auxiliary variables shown on flagged images only (02_kodeks.md: not validated).
AUX = ["people_count_band", "dress_items", "text_present", "legible_strings", "tourism_items", "memory_present"]
CORR_FIELDS = ["image_id", "variable", "model_value", "reviewed_value", "comment", "source_file", "ingested_at"]


def decisions_path(cfg: Config) -> Path:
    return cfg.dir("human") / "decisions.csv"


def read_decisions(cfg: Config) -> dict[str, str]:
    p = decisions_path(cfg)
    if not p.exists():
        raise KkaiError(f"Brak {p}. Zapisz decyzje z bramki 3 (kolumny variable, decision) na podstawie validation_table.csv.")
    out = {}
    for r in read_csv(p):
        if r["decision"] not in DECISIONS:
            raise KkaiError(f"decisions.csv: {r['variable']}: decyzja {r['decision']!r} spoza {DECISIONS}")
        out[r["variable"]] = r["decision"]
    missing = [v for v in SHEET_COLUMNS if v not in out]
    if missing:
        raise KkaiError(f"decisions.csv: brak decyzji dla {missing}")
    return out


def main_model_codes(cfg: Config) -> dict[str, dict]:
    rows = read_csv(cfg.data_dir / "ai" / "ai_codes.csv")
    runs = sorted({r["run_id"] for r in rows if r["set"] == "main"})
    if not runs:
        raise KkaiError("Brak kodów modelu dla etapu main. Najpierw: kkai code --set main")
    return {r["image_id"]: r for r in rows if r["run_id"] == runs[0]}


def validation_ids(cfg: Config) -> set[str]:
    p = cfg.dir("human") / "validation_ids.csv"
    return {r["image_id"] for r in read_csv(p)} if p.exists() else set()


def review_plan(cfg: Config) -> list[dict]:
    """Rows of the review sheet: image, variables to review, variables to code by hand."""
    dec = read_decisions(cfg)
    model = main_model_codes(cfg)
    val = validation_ids(cfg)
    main_ids = sorted(r["image_id"] for r in read_key(cfg) if r["stage"] == "main" and r["image_id"] not in val)
    plan = []
    for iid in main_ids:
        m = model.get(iid)
        flagged = m is None or m.get("valid") != "True" or m.get("needs_review_program") == "True"
        review = [v for v, d in dec.items() if d == "review_all" or (d == "review_flagged" and flagged)]
        manual = [v for v, d in dec.items() if d == "manual_all"]
        if m is None or m.get("valid") != "True":
            manual, review = manual + review, []  # no usable model codes: code by hand
        if review or manual:
            plan.append({"image_id": iid, "review": review, "manual": manual, "flagged": flagged, "model": m or {}})
    return plan


def run_review(cfg: Config) -> tuple[Path, Path]:
    human = cfg.dir("human")
    xlsx, page = human / "review_sheet.xlsx", human / "review.html"
    if xlsx.exists() or page.exists():
        raise KkaiError(f"{xlsx.name} albo {page.name} już istnieje. Nie nadpisuję.")
    plan = review_plan(cfg)
    variables = [v for v in SHEET_COLUMNS if any(v in p["review"] or v in p["manual"] for p in plan)]
    wb = Workbook()
    ws = wb.active
    ws.title = "review"
    header = ["image_id", "flagged", "review_reasons"]
    for v in variables:
        header += [f"model:{v}", f"corr:{v}"]
    header += [f"model:{a}" for a in AUX] + ["comment"]
    ws.append(header)
    grey = PatternFill("solid", fgColor="EEEEEE")
    for c in ws[1]:
        c.font = Font(bold=True)
    for p in plan:
        m = p["model"]
        row = [p["image_id"], p["flagged"], m.get("review_reasons", "")]
        for v in variables:
            if v in p["review"]:
                row += [m.get(v, ""), None]
            elif v in p["manual"]:
                row += ["(code by hand)", None]
            else:
                row += ["(no review)", "-"]
        row += [m.get(a, "") if p["flagged"] else "" for a in AUX] + [None]
        ws.append(row)
    for col in ws.iter_cols(min_row=1, max_row=ws.max_row):
        if str(col[0].value).startswith("model:"):
            for c in col[1:]:
                c.fill = grey
    help_ws = wb.create_sheet("help")
    for line in ["corr:<variable>: leave empty to accept the model code; type the correct code to change it.",
                 "'(code by hand)': the variable is coded manually here (gate 3 decision); the cell must not stay empty.",
                 "Multi-select columns: codes separated by ';'. Type '-' for none.",
                 "'(no review)' / '-': nothing to do for this variable in this image."]:
        help_ws.append([line])
    with open(xlsx, "xb") as f:
        wb.save(f)
    _review_html(cfg, page, plan, variables)
    n_manual = sum(len(p["manual"]) for p in plan)
    print(f"Arkusz przeglądu: {xlsx} ({len(plan)} obrazów, {n_manual} pól do kodowania ręcznego)\nStrona: {page}")
    return xlsx, page


def _review_html(cfg, page, plan, variables) -> None:
    items = []
    for p in plan:
        hits = list((cfg.data_dir / "blinded").glob(f"*/{p['image_id']}.png"))
        rel = os.path.relpath(hits[0], page.parent).replace(os.sep, "/") if hits else ""
        m = p["model"]
        cells = "".join(f"<tr><th>{v}</th><td>{html.escape(str(m.get(v, ''))) if v in p['review'] else '(ręcznie)'}</td></tr>"
                        for v in variables if v in p["review"] or v in p["manual"])
        aux = "".join(f"<tr><th>{a}</th><td>{html.escape(str(m.get(a, '')))}</td></tr>" for a in AUX) if p["flagged"] else ""
        items.append(f'<section id="{p["image_id"]}"><img src="{html.escape(rel)}" loading="lazy">'
                     f'<div><h2>{p["image_id"]}</h2><p>{html.escape(m.get("review_reasons", ""))}</p>'
                     f'<table>{cells}{aux}</table></div></section>')
    doc = f"""<!doctype html><html lang="pl"><head><meta charset="utf-8"><title>kkai: przegląd</title>
<style>body{{font-family:sans-serif;margin:16px}}section{{display:flex;gap:16px;margin-bottom:32px;flex-wrap:wrap}}
img{{width:640px;max-width:100%;border:1px solid #ccc}}td,th{{text-align:left;padding:2px 8px;font-size:14px}}</style>
</head><body>{''.join(items)}</body></html>"""
    write_new_text(page, doc)


def run_ingest_review(cfg: Config, sheet: str | Path) -> Path:
    """Read the filled review sheet, write review_corrections.csv and final_codes.csv."""
    sheet = Path(sheet)
    allowed = allowed_values(load_schema(str(cfg.path(cfg.raw["coder"]["schema"]))))
    plan = {p["image_id"]: p for p in review_plan(cfg)}
    ws = load_workbook(sheet, data_only=True)["review"]
    header = [c.value for c in ws[1]]
    errors, corrections = [], []
    now = utc_now()
    for rnum, vals in enumerate(ws.iter_rows(min_row=2, values_only=True), 2):
        rec = dict(zip(header, ["" if v is None else str(v).strip() for v in vals]))
        p = plan.get(rec["image_id"])
        if p is None:
            errors.append(f"wiersz {rnum}: {rec['image_id']} nie należy do przeglądu")
            continue
        for v in p["review"] + p["manual"]:
            val = rec.get(f"corr:{v}", "")
            # In a multi-select column '-' means "none"; an empty cell means "not coded".
            if v in p["manual"] and not val:
                errors.append(f"wiersz {rnum} ({rec['image_id']}): {v}: brak kodu (kodowanie ręczne)")
                continue
            if not val:
                continue
            multi = SHEET_COLUMNS[v][1]
            codes = [c.strip() for c in val.replace(",", ";").split(";") if c.strip()] if multi else [val]
            if val == "-" and multi:
                codes = []
            bad = [c for c in codes if c not in allowed[v]]
            if bad:
                errors.append(f"wiersz {rnum} ({rec['image_id']}): {v}: niedozwolone {bad}")
                continue
            corrections.append({"image_id": rec["image_id"], "variable": v,
                                "model_value": "" if v in p["manual"] else p["model"].get(v, ""),
                                "reviewed_value": ";".join(sorted(codes)), "comment": rec.get("comment", ""),
                                "source_file": f"{sheet.name} sha256={sha256_file(sheet)[:16]}", "ingested_at": now})
    if errors:
        for e in errors:
            print("  " + e)
        raise KkaiError(f"{len(errors)} problemów w arkuszu przeglądu. Nic nie zapisano.")
    human = cfg.dir("human")
    write_csv_once(human / "review_corrections.csv", corrections, CORR_FIELDS)
    path = build_final_codes(cfg, corrections)
    print(f"Poprawek: {len(corrections)}. Kody końcowe: {path}")
    return path


def build_final_codes(cfg: Config, corrections: list[dict]) -> Path:
    """Final code per variable and its source: human (validation set), review, manual or model."""
    dec = read_decisions(cfg)
    model = main_model_codes(cfg)
    corr = {(c["image_id"], c["variable"]): c["reviewed_value"] for c in corrections}
    human_rows = [r for r in read_csv(cfg.dir("human") / "human_codes.csv")
                  if r["set"] == "validation" and r["round"] == "1"]
    coder = cfg.raw.get("validation", {}).get("coder_id") or (human_rows[0]["coder_id"] if human_rows else None)
    human = {r["image_id"]: r for r in human_rows if r["coder_id"] == coder}
    rows = []
    for k in read_key(cfg):
        if k["stage"] != "main":
            continue
        iid = k["image_id"]
        row = {"image_id": iid}
        for v in SHEET_COLUMNS:
            if iid in human:
                val, src = human[iid].get(v, ""), "human_validation"
            elif (iid, v) in corr:
                val = corr[(iid, v)]
                src = "manual" if dec[v] == "manual_all" else ("review" if val != model.get(iid, {}).get(v) else "model_confirmed")
            elif dec[v] == "manual_all" or iid not in model or model[iid].get("valid") != "True":
                val, src = "", "missing"
            else:
                val, src = model[iid].get(v, ""), "model"
            row[v], row[f"{v}__source"] = val, src
        rows.append(row)
    fields = ["image_id"] + [f for v in SHEET_COLUMNS for f in (v, f"{v}__source")]
    path = cfg.dir("human") / "final_codes.csv"
    write_csv_once(path, sorted(rows, key=lambda r: r["image_id"]), fields)
    missing = sum(1 for r in rows for v in SHEET_COLUMNS if r[f"{v}__source"] == "missing")
    if missing:
        print(f"Uwaga: {missing} pól bez kodu końcowego (source = missing).")
    return path
