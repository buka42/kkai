"""Dry-run helper: fills coding sheets with plausible codes so that later steps can be tested.

Works only on data_dryrun/. Never used for real data.
"""

from __future__ import annotations

import random
import shutil
from pathlib import Path

from openpyxl import load_workbook

from .codebook import SHEET_COLUMNS, allowed_values, load_schema
from .config import Config
from .util import KkaiError, read_csv

NOISE = 0.1


def simulate_sheet(cfg: Config, sheet: str) -> Path:
    if not cfg.dry_run:
        raise KkaiError("simulate-human działa tylko w trybie --dry-run")
    human = cfg.dir("human")
    if sheet == "decisions":
        props = sorted((cfg.data_dir / "results" / "validation").glob("*/decisions_proposed.csv"))
        if not props:
            raise KkaiError("Brak decisions_proposed.csv. Najpierw: kkai validate --dry-run")
        shutil.copy(props[-1], human / "decisions.csv")
        print(f"Zapisano {human / 'decisions.csv'} (z {props[-1]})")
        return human / "decisions.csv"
    path = human / sheet
    if not path.exists():
        raise KkaiError(f"Brak {path}")
    allowed = allowed_values(load_schema(str(cfg.path(cfg.raw["coder"]["schema"]))))
    ai_path = cfg.data_dir / "ai" / "ai_codes.csv"
    model = {}
    if ai_path.exists():
        for r in read_csv(ai_path):
            if r["valid"] == "True":
                model.setdefault(r["image_id"], r)
    rng = random.Random(f"{cfg.seed}:simulate:{sheet}")
    wb = load_workbook(path)
    ws = wb[wb.sheetnames[0]]
    header = [c.value for c in ws[1]]
    for row in ws.iter_rows(min_row=2):
        iid = row[0].value
        for cell, name in zip(row, header):
            col = name.split(":", 1)[1] if name and name.startswith("corr:") else name
            if col not in SHEET_COLUMNS or (name.startswith("corr:") and cell.value == "-"):
                continue
            if name.startswith("corr:"):
                model_cell = row[header.index(f"model:{col}")].value
                if model_cell != "(code by hand)":
                    continue  # accept the model code
            multi = SHEET_COLUMNS[col][1]
            base = model.get(iid, {}).get(col)
            if base is None or rng.random() < NOISE:
                base = ";".join(rng.sample(allowed[col], rng.randint(0, 2))) if multi else rng.choice(allowed[col])
            if name.startswith("corr:") and multi and not base:
                base = "-"
            cell.value = base
        if "coder_id" in header:
            row[header.index("coder_id")].value = "SIM"
    out = path.with_name(path.stem + "_filled.xlsx")
    wb.save(out)
    print(f"Wypełniony arkusz (symulacja): {out}")
    return out
