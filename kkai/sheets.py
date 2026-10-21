"""Human coding sheets (XLSX with drop-down lists) and an HTML page with the images.

Sheets show only image_id: no condition, generator, prompt or model codes.
A sheet that already exists is never overwritten (it may hold your codes).
"""

from __future__ import annotations

import html
import os
import random
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

from .codebook import SECOND_CODER_COLUMNS, SHEET_COLUMNS, SHEET_FIELDS, allowed_values, load_schema
from .config import Config
from .prepare import read_key
from .util import KkaiError, read_csv, write_new_text

SHEET_SETS = ("pilot", "validation", "retest", "second")


def sheet_ids(cfg: Config, set_name: str) -> list[str]:
    human = cfg.dir("human")
    if set_name == "pilot":
        ids = sorted(r["image_id"] for r in read_key(cfg) if r["stage"] == "pilot")
        if not ids:
            raise KkaiError("Brak obrazów pilotażu. Najpierw: kkai prepare --stage pilot")
        random.Random(f"{cfg.seed}:sheet:pilot").shuffle(ids)
        return ids
    src = {"validation": "validation_ids.csv", "retest": "retest_ids.csv", "second": "second_ids.csv"}[set_name]
    path = human / src
    if not path.exists():
        raise KkaiError(f"Brak {path}. Najpierw: kkai sample")
    ids = [r["image_id"] for r in read_csv(path)]
    if set_name == "retest":
        # A different order than in the first round.
        random.Random(f"{cfg.seed}:sheet:retest").shuffle(ids)
    return ids


def sheet_fields(set_name: str) -> list[str]:
    if set_name == "second":
        return ["image_id", *SECOND_CODER_COLUMNS, "notes", "coder_id", "coded_at"]
    return list(SHEET_FIELDS)


def sheet_paths(cfg: Config, set_name: str) -> tuple[Path, Path]:
    human = cfg.dir("human")
    return human / f"{set_name}_sheet.xlsx", human / f"{set_name}_images.html"


def run_sheet(cfg: Config, set_name: str) -> tuple[Path, Path]:
    if set_name not in SHEET_SETS:
        raise KkaiError(f"Nieznany zbiór {set_name!r}; dozwolone: {', '.join(SHEET_SETS)}")
    xlsx, page = sheet_paths(cfg, set_name)
    if xlsx.exists() or page.exists():
        raise KkaiError(f"{xlsx.name} albo {page.name} już istnieje. Nie nadpisuję (arkusz może zawierać Twoje kody).")
    ids = sheet_ids(cfg, set_name)
    allowed = allowed_values(load_schema(str(cfg.path(cfg.raw["coder"]["schema"]))))
    fields = sheet_fields(set_name)
    _write_xlsx(xlsx, ids, fields, allowed)
    _write_html(cfg, page, ids, set_name)
    print(f"Arkusz: {xlsx}\nObrazy: {page}\nObrazów: {len(ids)}")
    return xlsx, page


def _write_xlsx(path: Path, ids: list[str], fields: list[str], allowed: dict) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "codes"
    lists = wb.create_sheet("lists")
    help_ws = wb.create_sheet("help")
    ws.append(fields)
    for c in ws[1]:
        c.font = Font(bold=True)
        c.fill = PatternFill("solid", fgColor="DDDDDD")
    for i in ids:
        ws.append([i] + [None] * (len(fields) - 1))
    ws.freeze_panes = "B2"
    n = len(ids) + 1

    help_ws.append(["column", "type", "allowed values"])
    for col_idx, name in enumerate(fields, 1):
        letter = get_column_letter(col_idx)
        ws.column_dimensions[letter].width = 16 if name != "notes" else 40
        if name not in SHEET_COLUMNS:
            continue
        values = allowed[name]
        multi = SHEET_COLUMNS[name][1]
        help_ws.append([name, "multi: codes separated by ';' (empty = none)" if multi else "single: drop-down",
                        ", ".join(values)])
        if multi:
            continue
        lcol = get_column_letter(lists.max_column + 1 if lists.max_row > 1 or lists["A1"].value else 1)
        lists[f"{lcol}1"] = name
        for r, v in enumerate(values, 2):
            lists[f"{lcol}{r}"] = v
        dv = DataValidation(type="list", formula1=f"=lists!${lcol}$2:${lcol}${len(values) + 1}",
                            allow_blank=True, showErrorMessage=True,
                            errorTitle="Niedozwolona wartość", error=f"Wybierz kod z listy ({name}).")
        ws.add_data_validation(dv)
        dv.add(f"{letter}2:{letter}{n}")
    lists.sheet_state = "hidden"
    path.parent.mkdir(parents=True, exist_ok=True)
    with open(path, "xb") as f:
        wb.save(f)


def _write_html(cfg: Config, page: Path, ids: list[str], set_name: str) -> None:
    blinded = cfg.data_dir / "blinded"
    items = []
    for i in ids:
        hits = list(blinded.glob(f"*/{i}.png"))
        if len(hits) != 1:
            raise KkaiError(f"Nie znajduję pliku {i}.png w blinded/")
        rel = os.path.relpath(hits[0], page.parent).replace(os.sep, "/")
        items.append(f'<figure id="{i}"><img src="{html.escape(rel)}" loading="lazy" alt="{i}">'
                     f"<figcaption>{i}</figcaption></figure>")
    doc = f"""<!doctype html>
<html lang="pl"><head><meta charset="utf-8"><title>kkai: {set_name}</title>
<style>
body {{ font-family: sans-serif; margin: 16px; background: #fafafa; }}
figure {{ margin: 0 0 32px; }}
img {{ max-width: 100%; width: 900px; display: block; border: 1px solid #ccc; }}
figcaption {{ font: 600 18px monospace; margin-top: 6px; }}
</style></head><body>
<p>{len(ids)} obrazów, kolejność jak w arkuszu.</p>
{chr(10).join(items)}
</body></html>
"""
    write_new_text(page, doc)
