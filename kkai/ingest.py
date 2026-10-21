"""Reading filled coding sheets into human_codes.csv (append only)."""

from __future__ import annotations

import csv
from pathlib import Path

from openpyxl import load_workbook

from .codebook import SHEET_COLUMNS, allowed_values, load_schema
from .config import Config
from .sheets import SHEET_SETS, sheet_fields, sheet_ids
from .util import KkaiError, read_csv, sha256_file, utc_now

HUMAN_FIELDS = ["image_id", "coder_id", "set", "round", *SHEET_COLUMNS, "notes", "coded_at",
                "source_file", "source_sha256", "ingested_at"]


def human_codes_path(cfg: Config) -> Path:
    return cfg.dir("human") / "human_codes.csv"


def check_sheet(cfg: Config, sheet: Path, set_name: str) -> tuple[list[dict], list[str]]:
    allowed = allowed_values(load_schema(str(cfg.path(cfg.raw["coder"]["schema"]))))
    fields = sheet_fields(set_name)
    wb = load_workbook(sheet, data_only=True)
    if "codes" not in wb.sheetnames:
        raise KkaiError(f"{sheet.name}: brak zakładki 'codes'")
    ws = wb["codes"]
    header = [c.value for c in ws[1]]
    missing_cols = [f for f in fields if f not in header]
    if missing_cols:
        raise KkaiError(f"{sheet.name}: brak kolumn {missing_cols}")
    expected_ids = set(sheet_ids(cfg, set_name))
    rows, errors, seen = [], [], set()
    for rnum, values in enumerate(ws.iter_rows(min_row=2, values_only=True), 2):
        rec = {h: ("" if v is None else str(v).strip()) for h, v in zip(header, values) if h}
        if not any(rec.values()):
            continue
        iid = rec.get("image_id", "")
        if iid not in expected_ids:
            errors.append(f"wiersz {rnum}: image_id {iid!r} nie należy do zbioru {set_name}")
        if iid in seen:
            errors.append(f"wiersz {rnum}: powtórzony image_id {iid}")
        seen.add(iid)
        for col in fields:
            if col not in SHEET_COLUMNS:
                continue
            val = rec.get(col, "")
            multi = SHEET_COLUMNS[col][1]
            if multi:
                codes = [c.strip() for c in val.replace(",", ";").split(";") if c.strip()]
                bad = [c for c in codes if c not in allowed[col]]
                if bad:
                    errors.append(f"wiersz {rnum} ({iid}): {col}: niedozwolone {bad}")
                rec[col] = ";".join(sorted(dict.fromkeys(codes)))
            else:
                if not val:
                    errors.append(f"wiersz {rnum} ({iid}): {col}: brak wartości")
                elif val not in allowed[col]:
                    errors.append(f"wiersz {rnum} ({iid}): {col}: niedozwolone {val!r}")
        if rec.get("people_count") == "none" and rec.get("roles"):
            errors.append(f"wiersz {rnum} ({iid}): people_count = none, a roles nie jest puste")
        rows.append(rec)
    for iid in sorted(expected_ids - seen):
        errors.append(f"brak wiersza dla {iid}")
    return rows, errors


def run_ingest(cfg: Config, sheet: str | Path, set_name: str, coder_id: str, round_no: int = 1) -> int:
    if set_name not in SHEET_SETS:
        raise KkaiError(f"Nieznany zbiór {set_name!r}; dozwolone: {', '.join(SHEET_SETS)}")
    sheet = Path(sheet)
    rows, errors = check_sheet(cfg, sheet, set_name)
    if errors:
        print(f"{sheet.name}: {len(errors)} problemów. Nic nie zapisano.")
        for e in errors:
            print("  " + e)
        raise KkaiError("Popraw arkusz i uruchom ponownie.")
    out = human_codes_path(cfg)
    existing = read_csv(out) if out.exists() else []
    key = lambda r: (r["image_id"], r["coder_id"], r["set"], str(r["round"]))
    done = {key(r) for r in existing}
    sha, now = sha256_file(sheet), utc_now()
    new = []
    for r in rows:
        rec = {f: r.get(f, "") for f in HUMAN_FIELDS}
        rec.update(coder_id=coder_id, set=set_name, round=round_no, source_file=sheet.name, source_sha256=sha,
                   ingested_at=now)
        if key(rec) in done:
            raise KkaiError(f"Kody {rec['image_id']} (koder {coder_id}, zbiór {set_name}, runda {round_no}) "
                            "są już wczytane. Nic nie zapisano.")
        new.append(rec)
    with open(out, "a", newline="", encoding="utf-8") as f:
        w = csv.DictWriter(f, fieldnames=HUMAN_FIELDS, lineterminator="\n")
        if not existing and f.tell() == 0:
            w.writeheader()
        w.writerows(new)
    print(f"Wczytano {len(new)} wierszy do {out}")
    return len(new)
